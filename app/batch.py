from __future__ import annotations

import asyncio
from io import BytesIO
from urllib.parse import urlparse

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.config import Settings
from app.domain import Brand, MaterialType
from app.reviewer import MaterialReviewer


HEADER_ALIASES = {
    "brand": ("品牌", "brand"),
    "material_type": ("物料类型", "material_type", "type"),
    "image_url": ("图片链接", "图片URL", "image_url", "url"),
}
OUTPUT_HEADERS = (
    "审核结果",
    "识别品牌",
    "识别物料类型",
    "置信度",
    "审核原因",
    "画质预警",
    "识别证据",
)


class BatchValidationError(ValueError):
    pass


def build_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "待审核图片"
    sheet.append(["品牌", "物料类型", "图片链接"])
    sheet.append(
        [
            "皇家",
            "灯箱",
            "https://static.51dh.com.cn/path/to/example.jpg",
        ]
    )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:C2"
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 72
    header_fill = PatternFill("solid", fgColor="1261A0")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _find_columns(sheet) -> dict[str, int]:
    headers = {
        str(cell.value).strip().lower(): cell.column
        for cell in sheet[1]
        if cell.value is not None
    }
    columns: dict[str, int] = {}
    for logical_name, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if alias.lower() in headers:
                columns[logical_name] = headers[alias.lower()]
                break
        if logical_name not in columns:
            raise BatchValidationError(f"缺少必填列：{aliases[0]}")
    return columns


def _allowed_host(url: str, allowed_hosts: tuple[str, ...]) -> bool:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        return False
    host = parsed.hostname.lower()
    return any(host == allowed or host.endswith(f".{allowed}") for allowed in allowed_hosts)


async def _download_image(
    client: httpx.AsyncClient, url: str, settings: Settings
) -> tuple[bytes, str]:
    if not _allowed_host(url, settings.allowed_image_hosts):
        raise BatchValidationError("图片域名不在 ALLOWED_IMAGE_HOSTS 白名单内")
    response = await client.get(url, follow_redirects=True)
    response.raise_for_status()
    final_url = str(response.url)
    if not _allowed_host(final_url, settings.allowed_image_hosts):
        raise BatchValidationError("重定向后的图片域名不在白名单内")
    payload = response.content
    if len(payload) > settings.max_image_bytes:
        raise BatchValidationError("图片超过大小限制")
    content_type = response.headers.get("content-type", "image/jpeg").split(";")[0]
    if not content_type.startswith("image/"):
        raise BatchValidationError("链接返回的不是图片")
    return payload, content_type


async def process_workbook(
    workbook_bytes: bytes,
    reviewer: MaterialReviewer,
    settings: Settings,
) -> bytes:
    try:
        workbook = load_workbook(BytesIO(workbook_bytes))
    except Exception as error:  # noqa: BLE001 - converted to user-facing validation
        raise BatchValidationError(f"无法读取 Excel：{error}") from error

    sheet = workbook.active
    columns = _find_columns(sheet)
    data_rows = [
        row
        for row in range(2, sheet.max_row + 1)
        if any(sheet.cell(row, columns[key]).value for key in columns)
    ]
    if not data_rows:
        raise BatchValidationError("Excel 中没有待审核数据")
    if len(data_rows) > settings.batch_max_rows:
        raise BatchValidationError(
            f"单次最多处理 {settings.batch_max_rows} 行，当前为 {len(data_rows)} 行"
        )

    output_start = sheet.max_column + 1
    for offset, header in enumerate(OUTPUT_HEADERS):
        cell = sheet.cell(1, output_start + offset, header)
        cell.fill = PatternFill("solid", fgColor="1261A0")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    semaphore = asyncio.Semaphore(settings.batch_concurrency)
    timeout = httpx.Timeout(settings.request_timeout_seconds)

    async with httpx.AsyncClient(timeout=timeout) as client:

        async def process_row(row_number: int) -> tuple[int, list[object]]:
            brand_value = str(sheet.cell(row_number, columns["brand"]).value or "").strip()
            type_value = str(
                sheet.cell(row_number, columns["material_type"]).value or ""
            ).strip()
            url = str(sheet.cell(row_number, columns["image_url"]).value or "").strip()
            try:
                brand = Brand(brand_value)
                material_type = MaterialType(type_value)
                if not url:
                    raise BatchValidationError("图片链接为空")
                async with semaphore:
                    payload, mime_type = await _download_image(client, url, settings)
                    result = await reviewer.review(
                        expected_brand=brand,
                        expected_material_type=material_type,
                        image_bytes=payload,
                        mime_type=mime_type,
                    )
                return row_number, [
                    {
                        "passed": "通过",
                        "rejected": "驳回",
                        "manual_review": "人工复核",
                    }[result.status],
                    result.detected_brand.value if result.detected_brand else "",
                    (
                        result.detected_material_type.value
                        if result.detected_material_type
                        else ""
                    ),
                    result.confidence,
                    "；".join(result.reasons),
                    "；".join(result.quality.warnings),
                    "；".join(result.evidence),
                ]
            except Exception as error:  # noqa: BLE001 - row-level errors belong in output
                return row_number, ["处理失败", "", "", 0, str(error), "", ""]

        results = await asyncio.gather(*(process_row(row) for row in data_rows))

    for row_number, values in results:
        for offset, value in enumerate(values):
            sheet.cell(row_number, output_start + offset, value)
        sheet.cell(row_number, output_start + 3).number_format = "0.0%"

    widths = [14, 14, 24, 12, 54, 36, 54]
    for offset, width in enumerate(widths):
        sheet.column_dimensions[
            sheet.cell(1, output_start + offset).column_letter
        ].width = width
    sheet.freeze_panes = sheet.freeze_panes or "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()

