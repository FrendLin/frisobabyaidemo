from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

import app.batch as batch_module
from app.batch import process_workbook
from app.config import Settings
from app.domain import Brand, MaterialType
from app.models import DetectedMaterial, ExternalImageQuality
from app.reviewer import MaterialReviewer
from conftest import FakeProvider


@pytest.mark.asyncio
async def test_batch_appends_results(monkeypatch, image_bytes: bytes) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["品牌", "物料类型", "图片链接"])
    sheet.append(["皇家", "灯箱", "https://static.51dh.com.cn/example.jpg"])
    source = BytesIO()
    workbook.save(source)

    async def fake_download(client, url, settings):
        del client, url, settings
        return image_bytes, "image/jpeg"

    monkeypatch.setattr(batch_module, "_download_image", fake_download)
    provider = FakeProvider(
        DetectedMaterial(
            brand=Brand.ROYAL,
            material_type=MaterialType.LIGHTBOX,
            confidence=0.91,
        )
    )
    reviewer = MaterialReviewer(provider, Settings())
    output = await process_workbook(source.getvalue(), reviewer, Settings())

    result = load_workbook(BytesIO(output), data_only=True).active
    headers = [cell.value for cell in result[1]]
    assert "审核结果" in headers
    assert result.cell(2, headers.index("审核结果") + 1).value == "通过"


@pytest.mark.asyncio
async def test_batch_quality_check_reuses_source_url(
    monkeypatch,
    image_bytes: bytes,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["品牌", "物料类型", "图片链接"])
    source_url = "https://static.51dh.com.cn/example.jpg"
    sheet.append(["皇家", "灯箱", source_url])
    source = BytesIO()
    workbook.save(source)

    async def fake_download(client, url, settings):
        del client, url, settings
        return image_bytes, "image/jpeg"

    monkeypatch.setattr(batch_module, "_download_image", fake_download)
    provider = FakeProvider(
        DetectedMaterial(
            brand=Brand.ROYAL,
            material_type=MaterialType.LIGHTBOX,
            confidence=0.91,
        )
    )

    class FakeQualityChecker:
        called_url = ""

        async def inspect(self, image_url: str) -> ExternalImageQuality:
            self.called_url = image_url
            return ExternalImageQuality(acceptable=True, description="低")

    class FailingPublisher:
        async def publish(self, payload: bytes, mime_type: str) -> str:
            raise AssertionError("批量图片已有 URL，不应上传 OSS")

    checker = FakeQualityChecker()
    reviewer = MaterialReviewer(
        provider,
        Settings(),
        quality_checker=checker,
        image_publisher=FailingPublisher(),
    )
    output = await process_workbook(
        source.getvalue(),
        reviewer,
        Settings(),
        check_image_quality=True,
    )

    result = load_workbook(BytesIO(output), data_only=True).active
    headers = [cell.value for cell in result[1]]
    assert result.cell(2, headers.index("审核结果") + 1).value == "通过"
    assert checker.called_url == source_url
