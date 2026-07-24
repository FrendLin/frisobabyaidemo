#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import math
import random
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_SEED = 20260724
DEFAULT_TEST_COUNT = 44


def read_samples(source: Path) -> list[dict[str, str]]:
    workbook = load_workbook(source, data_only=True, read_only=True)
    if "图片链接" not in workbook.sheetnames:
        raise ValueError("工作簿缺少“图片链接”工作表")
    samples = []
    for row_number, values in enumerate(
        workbook["图片链接"].iter_rows(min_row=2, values_only=True), start=2
    ):
        brand, material_type, image_url = values[:3]
        if not brand or not material_type or not image_url:
            continue
        samples.append(
            {
                "sample_id": f"MS-{row_number - 1:04d}",
                "brand": str(brand).strip(),
                "material_type": str(material_type).strip(),
                "image_url": str(image_url).strip(),
            }
        )
    return samples


def assign_stratified_split(
    samples: list[dict[str, str]], test_count: int, seed: int
) -> None:
    if not 0 < test_count < len(samples):
        raise ValueError("test_count 必须位于 1 和样本总数之间")

    groups: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for sample in samples:
        groups[(sample["brand"], sample["material_type"])].append(sample)

    randomizer = random.Random(seed)
    for group in groups.values():
        randomizer.shuffle(group)

    allocations: dict[tuple[str, str], int] = {}
    remainders = []
    allocated = 0
    for label, group in groups.items():
        exact = len(group) * test_count / len(samples)
        minimum = 1 if len(group) >= 2 else 0
        maximum = len(group) - 1
        count = min(maximum, max(minimum, math.floor(exact)))
        allocations[label] = count
        allocated += count
        remainders.append((exact - math.floor(exact), len(group), label))

    for _, _, label in sorted(remainders, reverse=True):
        if allocated == test_count:
            break
        if allocations[label] < len(groups[label]) - 1:
            allocations[label] += 1
            allocated += 1

    if allocated != test_count:
        raise RuntimeError(f"无法分配 {test_count} 条测试样本，实际分配 {allocated}")

    for label, group in groups.items():
        for index, sample in enumerate(group):
            sample["split"] = "test" if index < allocations[label] else "train"


def main() -> None:
    parser = argparse.ArgumentParser(description="生成可复现的 80/20 分层数据清单")
    parser.add_argument("source", type=Path, help="客户提供的《大型品牌物料审核》Excel")
    parser.add_argument(
        "--output", type=Path, default=Path("data/dataset_manifest.csv")
    )
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    parser.add_argument("--test-count", type=int, default=DEFAULT_TEST_COUNT)
    args = parser.parse_args()

    samples = read_samples(args.source)
    assign_stratified_split(samples, args.test_count, args.seed)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=("sample_id", "brand", "material_type", "image_url", "split"),
        )
        writer.writeheader()
        writer.writerows(sorted(samples, key=lambda item: item["sample_id"]))

    train_count = sum(sample["split"] == "train" for sample in samples)
    test_count = sum(sample["split"] == "test" for sample in samples)
    print(
        f"已写入 {args.output}：总计 {len(samples)}，"
        f"训练 {train_count}，测试 {test_count}，seed={args.seed}"
    )


if __name__ == "__main__":
    main()

