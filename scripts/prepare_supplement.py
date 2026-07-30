#!/usr/bin/env python3
"""把客户补充的《品牌物料补充训练》Excel 接入版本化数据准备流程。

用法：
    python scripts/prepare_supplement.py <补充Excel> \\
        --version 20260727 \\
        --existing-manifest data/dataset_manifest.csv \\
        --output data/dataset_supplement_20260727.csv

流程：读取 Excel → 规范化同义标签（包柱→包柱画面）→ 校验品牌/类型合法性 →
对内去重并排除与既有清单重复的 URL → 写入带版本号的独立补充清单 →
打印按品牌、物料类型、联合标签的统计与非法/缺失行报告。历史基线清单不被修改。
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from app.dataset import (
    load_existing_urls,
    prepare_supplement,
    write_supplement_manifest,
)


def read_rows(source: Path) -> list[tuple[str, str, str]]:
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header = [str(c.value).strip() if c.value else "" for c in next(sheet.iter_rows(max_row=1))]
    # 兼容“月份/品牌/物料类型/图片”与“品牌/物料类型/图片链接”两种表头。
    def col(*names: str) -> int | None:
        for name in names:
            if name in header:
                return header.index(name)
        return None

    brand_idx = col("品牌", "brand")
    type_idx = col("物料类型", "material_type", "type")
    url_idx = col("图片", "图片链接", "图片URL", "image_url", "url")
    if brand_idx is None or type_idx is None or url_idx is None:
        raise SystemExit(f"表头缺少必要列，实际表头：{header}")

    rows: list[tuple[str, str, str]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        brand = values[brand_idx] if brand_idx < len(values) else None
        mtype = values[type_idx] if type_idx < len(values) else None
        url = values[url_idx] if url_idx < len(values) else None
        if not any((brand, mtype, url)):
            continue
        rows.append((str(brand or ""), str(mtype or ""), str(url or "")))
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="版本化接入补充训练数据")
    parser.add_argument("source", type=Path, help="客户补充的 Excel 文件")
    parser.add_argument("--version", required=True, help="数据版本号，如 20260727")
    parser.add_argument(
        "--existing-manifest",
        type=Path,
        default=Path("data/dataset_manifest.csv"),
        help="历史基线清单，用于 URL 去重（不会被修改）",
    )
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--id-prefix", default="SUP")
    args = parser.parse_args()

    output = args.output or Path(f"data/dataset_supplement_{args.version}.csv")
    rows = read_rows(args.source)
    existing = load_existing_urls(args.existing_manifest)
    result = prepare_supplement(
        rows,
        version=args.version,
        existing_urls=existing,
        id_prefix=args.id_prefix,
    )
    write_supplement_manifest(output, result)

    print(f"数据版本：{result.version}")
    print(f"原始行数：{len(rows)}")
    print(f"历史基线 URL 数：{len(existing)}")
    print(f"有效补充样本：{len(result.samples)}（已写入 {output}）")
    print(
        f"规范化标签：{result.normalized} 行；"
        f"补充集内重复：{result.duplicates_within}；"
        f"与既有清单重复：{result.duplicates_existing}"
    )
    print(f"非法/缺失行：{len(result.errors)}")
    for err in result.errors:
        print(f"  第 {err.row_number} 行：{err.reason}｜{err.brand}/{err.material_type}/{err.image_url}")

    print("\n按品牌统计：")
    for brand, count in sorted(result.brand_counts().items()):
        print(f"  {brand}: {count}")
    print("按物料类型统计：")
    for mtype, count in sorted(result.material_counts().items()):
        print(f"  {mtype}: {count}")
    print("按联合标签统计：")
    for (brand, mtype), count in sorted(result.joint_counts().items()):
        print(f"  {brand} · {mtype}: {count}")


if __name__ == "__main__":
    main()
